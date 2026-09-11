from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from docx import Document

_TIPO_POR_PREFIXO: dict[str, str] = {
    "01_": "termo_abertura",
    "02_": "cronograma",
    "03_": "mapa_beneficios",
    "04_": "riscos_problemas",
    "05_": "mudancas",
    "06_": "relatorio_encerramento",
}


@dataclass(frozen=True)
class TextoExtraido:
    texto: str
    projeto_id: str
    tipo_documento: str
    arquivo_origem: str
    secao: str = ""
    metadados_extras: dict = field(default_factory=dict)


def _projeto_id(caminho: Path) -> str:
    m = re.match(r"(SYN-\d+)", caminho.parent.name)
    return m.group(1) if m else "PORTFOLIO"


def _tipo_documento(caminho: Path) -> str:
    nome = caminho.name
    for prefixo, tipo in _TIPO_POR_PREFIXO.items():
        if nome.startswith(prefixo):
            return tipo
    if "Portfolio" in nome or "portfolio" in nome:
        return "portfolio"
    return "desconhecido"


def _nome_de_origem(caminho: Path) -> str:
    """O NOME do arquivo, nunca o caminho.

    Aqui ficava `str(caminho)`, o caminho absoluto da máquina de quem rodou a
    indexação. A coleção atual guarda coisas como
    "/Users/<alguém>/Downloads/base_sintetica_metro/SYN-04.../02_Cronograma.xlsx".

    Duas consequências, e a segunda é irreversível:

      1. O RNF12 pede referência RECUPERÁVEL, e um caminho no computador de
         outra pessoa não é recuperável por mais ninguém. O campo aparece na
         resposta da API e na interface.
      2. O valor vai para `auditoria.mensagem_fonte.arquivo_origem`, que é NOT
         NULL e está sob `REVOKE UPDATE, DELETE`. Errado ali, fica errado para
         sempre.

    O nome sozinho não é ambíguo porque `projeto_id` viaja ao lado: dois
    projetos têm "02_Cronograma.xlsx", mas o par (projeto_id, arquivo_origem)
    distingue. Os dois são gravados nas duas pontas, no índice e na trilha.
    """
    return caminho.name


def extrair_docx(caminho: Path) -> list[TextoExtraido]:
    doc = Document(caminho)
    projeto_id = _projeto_id(caminho)
    tipo = _tipo_documento(caminho)
    arquivo = _nome_de_origem(caminho)

    resultado: list[TextoExtraido] = []
    secao_atual = ""

    for paragrafo in doc.paragraphs:
        texto = paragrafo.text.strip()
        if not texto:
            continue

        estilo = paragrafo.style.name if paragrafo.style else ""
        if "Heading" in estilo or "heading" in estilo:
            secao_atual = texto
            continue

        resultado.append(TextoExtraido(
            texto=texto,
            projeto_id=projeto_id,
            tipo_documento=tipo,
            arquivo_origem=arquivo,
            secao=secao_atual,
        ))

    return resultado


def _linha_para_texto(cabecalhos: list[str], linha: tuple) -> str:
    partes: list[str] = []
    for cabecalho, celula in zip(cabecalhos, linha):
        valor = celula.value
        if valor is None or str(valor).strip() == "":
            continue
        partes.append(f"{cabecalho}: {valor}")
    return " | ".join(partes)


def extrair_xlsx(caminho: Path) -> list[TextoExtraido]:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    projeto_id = _projeto_id(caminho)
    tipo = _tipo_documento(caminho)
    arquivo = _nome_de_origem(caminho)

    resultado: list[TextoExtraido] = []

    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows())
        if not rows:
            continue

        # Primeira linha com pelo menos 2 células não-vazias é o cabeçalho
        idx_cabecalho = 0
        for i, row in enumerate(rows):
            valores = [c.value for c in row if c.value is not None]
            if len(valores) >= 2:
                idx_cabecalho = i
                break

        cabecalhos = [
            str(c.value).strip() if c.value else f"Col{j}"
            for j, c in enumerate(rows[idx_cabecalho])
        ]

        for row in rows[idx_cabecalho + 1:]:
            texto = _linha_para_texto(cabecalhos, row)
            if not texto or len(texto) < 10:
                continue
            resultado.append(TextoExtraido(
                texto=texto,
                projeto_id=projeto_id,
                tipo_documento=tipo,
                arquivo_origem=arquivo,
                secao=sheet.title,
            ))

    return resultado


def extrair(caminho: Path) -> list[TextoExtraido]:
    sufixo = caminho.suffix.lower()
    if sufixo == ".docx":
        return extrair_docx(caminho)
    if sufixo == ".xlsx":
        return extrair_xlsx(caminho)
    return []
