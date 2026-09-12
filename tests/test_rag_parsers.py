# Testes dos parsers de DOCX e XLSX.
#
# O módulo não tinha teste nenhum, e o que ele produz é permanente: os
# metadados que saem daqui vão para o índice vetorial e, a partir do repositório
# de conversas, para `auditoria.mensagem_fonte` — que é NOT NULL e está sob
# `REVOKE UPDATE, DELETE`. Metadado errado gravado na trilha não se corrige.
#
# Os testes montam arquivos de verdade, em vez de simular `python-docx` e
# `openpyxl`. As regras aqui (qual linha é cabeçalho, qual parágrafo é título)
# dependem do comportamento real dessas bibliotecas, e um duplo só provaria que
# o duplo concorda com o teste.

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document

from rag.parsers import (
    TextoExtraido,
    _nome_de_origem,
    _projeto_id,
    _tipo_documento,
    extrair,
)


def _criar_docx(caminho: Path, blocos: list[tuple[str, str]]) -> None:
    """`blocos` é uma lista de (estilo, texto); estilo "" vira parágrafo comum."""
    doc = Document()
    for estilo, texto in blocos:
        if estilo:
            doc.add_paragraph(texto, style=estilo)
        else:
            doc.add_paragraph(texto)
    doc.save(caminho)


def _criar_xlsx(caminho: Path, linhas: list[list], titulo_da_aba: str = "Planilha") -> None:
    wb = openpyxl.Workbook()
    aba = wb.active
    aba.title = titulo_da_aba
    for linha in linhas:
        aba.append(linha)
    wb.save(caminho)


class _ComPastaTemporaria(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.base = Path(self._tmp.name)

    def pasta_de_projeto(self, nome: str) -> Path:
        pasta = self.base / nome
        pasta.mkdir(parents=True, exist_ok=True)
        return pasta


class TesteOrigemDoArquivo(_ComPastaTemporaria):
    # ESTE É O TESTE DE UMA CORREÇÃO, e a razão dele é o que já está gravado na
    # base: a coleção em produção guarda
    # "/Users/<alguém>/Downloads/base_sintetica_metro/..." em `arquivo_origem`,
    # porque aqui havia `str(caminho)`.
    def test_guarda_o_nome_e_nao_o_caminho(self):
        self.assertEqual(
            _nome_de_origem(Path("/Users/alguem/Downloads/base/SYN-04/02_Cronograma.xlsx")),
            "02_Cronograma.xlsx",
        )

    def test_nao_vaza_a_arvore_de_diretorios(self):
        nome = _nome_de_origem(Path("/Users/alguem/Downloads/base/SYN-04/02_Cronograma.xlsx"))
        self.assertNotIn("/", nome)
        self.assertNotIn("Users", nome)

    def test_o_extrator_usa_o_nome(self):
        pasta = self.pasta_de_projeto("SYN-04_integracao")
        caminho = pasta / "01_Termo_de_Abertura.docx"
        _criar_docx(caminho, [("", "O objetivo do projeto é integrar a comunicação.")])

        extraidos = extrair(caminho)
        self.assertTrue(extraidos)
        for item in extraidos:
            self.assertEqual(item.arquivo_origem, "01_Termo_de_Abertura.docx")


class TesteIdentificacaoDoProjeto(unittest.TestCase):
    def test_le_o_codigo_da_pasta(self):
        self.assertEqual(_projeto_id(Path("/b/SYN-04_integracao/02_Cronograma.xlsx")), "SYN-04")

    def test_pasta_sem_codigo_vira_portfolio(self):
        # A planilha de portfólio não pertence a projeto algum.
        self.assertEqual(_projeto_id(Path("/b/portfolio/Portfolio.xlsx")), "PORTFOLIO")

    def test_codigo_precisa_estar_no_inicio(self):
        self.assertEqual(_projeto_id(Path("/b/backup_SYN-04/x.xlsx")), "PORTFOLIO")


class TesteTipoDeDocumento(unittest.TestCase):
    def test_tipos_pelo_prefixo_numerico(self):
        casos = {
            "01_Termo_de_Abertura.docx": "termo_abertura",
            "02_Cronograma.xlsx": "cronograma",
            "03_Mapa_de_Beneficios.xlsx": "mapa_beneficios",
            "04_Riscos_e_Problemas.xlsx": "riscos_problemas",
            "05_Mudancas.xlsx": "mudancas",
            "06_Relatorio_Anual_Encerramento.docx": "relatorio_encerramento",
        }
        for nome, esperado in casos.items():
            with self.subTest(nome=nome):
                self.assertEqual(_tipo_documento(Path("/b/SYN-01/" + nome)), esperado)

    def test_portfolio_pelo_nome(self):
        self.assertEqual(_tipo_documento(Path("/b/Portfolio_Sintetico_2026.xlsx")), "portfolio")

    def test_fora_da_convencao_vira_desconhecido(self):
        # Não é hipótese: a coleção atual tem 21 chunks assim, vindos de uma
        # pasta que não segue a nomenclatura. Eles são indexados e continuam
        # buscáveis, só sem poder ser filtrados por tipo.
        self.assertEqual(_tipo_documento(Path("/b/VETORIZE/Teste Prompt.docx")), "desconhecido")


class TesteExtracaoDeDocx(_ComPastaTemporaria):
    def test_titulo_vira_secao_dos_paragrafos_seguintes(self):
        pasta = self.pasta_de_projeto("SYN-01_ventilacao")
        caminho = pasta / "01_Termo_de_Abertura.docx"
        _criar_docx(
            caminho,
            [
                ("Heading 1", "Justificativa"),
                ("", "A ventilação atual opera acima da vida útil prevista."),
                ("Heading 1", "Escopo"),
                ("", "Substituição dos ventiladores das estações da linha."),
            ],
        )

        extraidos = extrair(caminho)
        secoes = {e.secao for e in extraidos}
        self.assertEqual(secoes, {"Justificativa", "Escopo"})

    def test_o_titulo_nao_vira_conteudo(self):
        # Ele é rótulo de seção; indexá-lo como texto criaria um chunk sem
        # informação competindo na busca.
        pasta = self.pasta_de_projeto("SYN-01_ventilacao")
        caminho = pasta / "01_Termo_de_Abertura.docx"
        _criar_docx(caminho, [("Heading 1", "Justificativa"), ("", "Texto do corpo aqui.")])

        textos = [e.texto for e in extrair(caminho)]
        self.assertEqual(textos, ["Texto do corpo aqui."])

    def test_paragrafo_vazio_e_descartado(self):
        pasta = self.pasta_de_projeto("SYN-01_ventilacao")
        caminho = pasta / "01_Termo_de_Abertura.docx"
        _criar_docx(caminho, [("", "Primeiro."), ("", "   "), ("", "Segundo.")])

        self.assertEqual([e.texto for e in extrair(caminho)], ["Primeiro.", "Segundo."])

    def test_metadados_acompanham_cada_trecho(self):
        pasta = self.pasta_de_projeto("SYN-07_manutencao")
        caminho = pasta / "01_Termo_de_Abertura.docx"
        _criar_docx(caminho, [("", "Conteúdo suficiente para o trecho.")])

        item = extrair(caminho)[0]
        self.assertEqual(item.projeto_id, "SYN-07")
        self.assertEqual(item.tipo_documento, "termo_abertura")


class TesteExtracaoDeXlsx(_ComPastaTemporaria):
    def test_cabecalho_rotula_cada_celula(self):
        pasta = self.pasta_de_projeto("SYN-04_integracao")
        caminho = pasta / "02_Cronograma.xlsx"
        _criar_xlsx(
            caminho,
            [["Marco", "Data"], ["Entrega do piloto", "2026-10-15"]],
            titulo_da_aba="Marcos",
        )

        item = extrair(caminho)[0]
        self.assertIn("Marco: Entrega do piloto", item.texto)
        self.assertIn("Data: 2026-10-15", item.texto)
        self.assertEqual(item.secao, "Marcos")

    def test_celula_vazia_nao_vira_rotulo_solto(self):
        pasta = self.pasta_de_projeto("SYN-04_integracao")
        caminho = pasta / "02_Cronograma.xlsx"
        _criar_xlsx(caminho, [["Marco", "Data"], ["Entrega do piloto muito importante", None]])

        self.assertNotIn("Data:", extrair(caminho)[0].texto)

    def test_linha_curta_demais_e_descartada(self):
        # Linhas quase vazias são ruído de planilha e produziriam chunks sem
        # conteúdo recuperável.
        #
        # O limiar de 10 caracteres mede o texto JÁ ROTULADO, cabeçalho
        # incluído — e não o conteúdo da célula. A consequência é que o filtro
        # afrouxa quando os cabeçalhos são longos: com "Marco"/"Data" a linha
        # ["x", "y"] vira "Marco: x | Data: y" e sobrevive. É por isso que o
        # caso abaixo tem uma célula só.
        pasta = self.pasta_de_projeto("SYN-04_integracao")
        caminho = pasta / "02_Cronograma.xlsx"
        _criar_xlsx(
            caminho,
            [["Marco", "Data"], ["x", None], ["Entrega do piloto", "2026-10-15"]],
        )

        extraidos = extrair(caminho)
        self.assertEqual(len(extraidos), 1)
        self.assertIn("Entrega do piloto", extraidos[0].texto)

    def test_cabecalho_e_a_primeira_linha_com_duas_celulas(self):
        # Planilhas reais costumam ter título numa linha só antes do cabeçalho.
        pasta = self.pasta_de_projeto("SYN-04_integracao")
        caminho = pasta / "02_Cronograma.xlsx"
        _criar_xlsx(
            caminho,
            [["Cronograma do projeto"], ["Marco", "Data"], ["Entrega do piloto", "2026-10-15"]],
        )

        self.assertIn("Marco: Entrega do piloto", extrair(caminho)[0].texto)


class TesteDespacho(_ComPastaTemporaria):
    def test_extensao_desconhecida_devolve_vazio_em_vez_de_erro(self):
        # `indexar_pasta` varre a pasta inteira; um PDF solto não pode derrubar
        # a indexação dos outros arquivos.
        caminho = self.base / "leiame.txt"
        caminho.write_text("nada", encoding="utf-8")
        self.assertEqual(extrair(caminho), [])

    def test_extensao_em_maiuscula_e_aceita(self):
        pasta = self.pasta_de_projeto("SYN-01_ventilacao")
        caminho = pasta / "01_Termo_de_Abertura.DOCX"
        _criar_docx(caminho, [("", "Conteúdo suficiente para o trecho.")])
        self.assertTrue(extrair(caminho))

    def test_devolve_textoextraido(self):
        pasta = self.pasta_de_projeto("SYN-01_ventilacao")
        caminho = pasta / "01_Termo_de_Abertura.docx"
        _criar_docx(caminho, [("", "Conteúdo suficiente para o trecho.")])
        self.assertIsInstance(extrair(caminho)[0], TextoExtraido)


if __name__ == "__main__":
    unittest.main(verbosity=2)
