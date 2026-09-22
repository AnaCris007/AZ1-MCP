"""RAG: busca, contrato, erros e ingestão — casos TI-55 a TI-58 (Seção 6.4.4).

A cadeia do RAG tem três elos, e eles têm custos de execução muito diferentes.
Esta suíte separa o que dá para verificar sempre do que exige infraestrutura e
chamada paga, em vez de pular tudo junto:

- **Ingestão (TI-58)** — extração e fragmentação são funções puras sobre
  arquivos. Rodam sempre, sem banco e sem provedor. É onde estão as lacunas
  mais concretas do pipeline, e elas ficam registradas como caracterização.
- **Contrato da rota (TI-56)** — os limites de `n_resultados`, o campo ausente
  e a lista vazia são decididos pelo schema e pela rota, antes de qualquer
  busca. Rodam sempre.
- **Erros (TI-57)** — a falha é injetada no elo indicado, e o que se verifica é
  a tradução para o contrato HTTP. Rodam sempre.
- **Busca de ponta a ponta (TI-55)** — precisa de `vecs` sobre PostgreSQL E de
  embeddings de documento, que não estão gravados em fita (a fita de embedding
  é de CONSULTA, tarefa diferente). Indexar cobra chamada real ao provedor, e
  por isso este bloco só roda sob `TEST_RAG_DB_URL` explícito.

Execução:

    python -m unittest tests.test_integracao_rag -v

Para incluir TI-55, que gasta chamadas reais de embedding:

    export TEST_RAG_DB_URL=postgresql://az1:az1@127.0.0.1:5432/az1_teste
"""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document

from az1_api.dependencies import get_document_searcher
from rag.chunker import chunkar
from rag.parsers import extrair, extrair_docx, extrair_xlsx
from rag.retriever import ResultadoBusca
from tests.apoio_integracao import cliente, limpar_overrides

RAG_DSN = os.environ.get("TEST_RAG_DB_URL", "")

MOTIVO_SEM_BANCO = (
    "TEST_RAG_DB_URL não definido. Este bloco indexa em `vecs` sobre PostgreSQL e "
    "gasta chamadas reais de embedding; precisa de opt-in explícito."
)


def _docx_com_paragrafos_e_tabela(destino: Path) -> Path:
    """DOCX sintético: um título, dois parágrafos e uma tabela."""
    documento = Document()
    documento.add_heading("Escopo do empreendimento", level=1)
    documento.add_paragraph("O projeto contempla a expansão de quatro estações.")
    documento.add_paragraph("O prazo contratual previsto é de trinta e seis meses.")

    tabela = documento.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text = "Marco"
    tabela.cell(0, 1).text = "Data"
    tabela.cell(1, 0).text = "Conclusão da via permanente"
    tabela.cell(1, 1).text = "2027-03-31"

    caminho = destino / "01_TAP.docx"
    documento.save(caminho)
    return caminho


def _xlsx_com_abas_e_vazios(destino: Path) -> Path:
    """XLSX sintético: duas abas, uma delas com células vazias."""
    planilha = openpyxl.Workbook()
    riscos = planilha.active
    riscos.title = "Riscos"
    riscos.append(["Código", "Descrição", "Severidade"])
    riscos.append(["R-01", "Atraso na desapropriação da área central", "alta"])
    riscos.append([None, None, None])
    riscos.append(["R-02", "Interferência com rede de água potável", "media"])

    problemas = planilha.create_sheet("Problemas")
    problemas.append(["Código", "Descrição"])
    problemas.append(["P-01", "Equipamento de perfuração indisponível no trecho norte"])

    caminho = destino / "04_Riscos.xlsx"
    planilha.save(caminho)
    return caminho


class _BuscadorControlado:
    """Substitui `rag.retriever.buscar` e guarda como foi chamado.

    O que ele NÃO faz é fingir uma busca semântica: os resultados são fixos. O
    que TI-56 e TI-57 verificam é o contrato em volta da busca — validação,
    repasse dos filtros e tradução de erro —, e isso não depende de o vetor ter
    sido calculado.
    """

    def __init__(self, *, resultados: list[ResultadoBusca] | None = None, erro: Exception | None = None) -> None:
        self.chamadas: list[dict] = []
        self._resultados = resultados if resultados is not None else []
        self._erro = erro

    def __call__(self, query: str, **argumentos: object) -> list[ResultadoBusca]:
        self.chamadas.append({"query": query, **argumentos})
        if self._erro is not None:
            raise self._erro
        return self._resultados


def _resultado(**trocas) -> ResultadoBusca:
    base = {
        "texto": "O avanço físico registrado no período é de 62%.",
        "score": 0.81,
        "projeto_id": "SYN-04",
        "tipo_documento": "cronograma",
        "secao": "Marcos",
        "arquivo_origem": "02_Cronograma.xlsx",
        "chunk_id": "chunk-abc",
    }
    base.update(trocas)
    return ResultadoBusca(**base)


class TestContratoDaBuscaRag(unittest.TestCase):
    """TI-56: limites, campo ausente e lista vazia."""

    def tearDown(self) -> None:
        limpar_overrides()

    def _buscar(self, corpo: dict, buscador: _BuscadorControlado | None = None):
        buscador = buscador or _BuscadorControlado()
        http = cliente({get_document_searcher: lambda: buscador})
        return http.post("/api/v1/rag/search", json=corpo), buscador

    def test_limites_invalidos_e_campo_ausente_sao_recusados(self) -> None:
        """`n_resultados` fora de 1..20 e `query` ausente param no schema."""
        casos = {
            "n_resultados abaixo do mínimo": {"query": "avanço", "n_resultados": 0},
            "n_resultados acima do máximo": {"query": "avanço", "n_resultados": 21},
            "campo query ausente": {"n_resultados": 5},
        }

        for descricao, corpo in casos.items():
            with self.subTest(caso=descricao):
                resposta, buscador = self._buscar(corpo)

                self.assertEqual(resposta.status_code, 422)
                self.assertEqual(buscador.chamadas, [], "a busca foi acionada com entrada inválida")

    def test_limites_validos_nas_bordas_sao_aceitos(self) -> None:
        """1 e 20 passam. O contraste que dá sentido ao caso acima."""
        for n in (1, 20):
            with self.subTest(n_resultados=n):
                resposta, buscador = self._buscar({"query": "avanço", "n_resultados": n})

                self.assertEqual(resposta.status_code, 200)
                self.assertEqual(buscador.chamadas[0]["n_resultados"], n)

    def test_filtro_sem_correspondencia_devolve_lista_vazia_com_200(self) -> None:
        """Nenhum registro correspondente não é erro: é `200` com `resultados=[]`."""
        resposta, buscador = self._buscar(
            {"query": "assunto inexistente", "projeto_id": "SYN-99", "tipo_documento": "mudancas"},
            _BuscadorControlado(resultados=[]),
        )

        self.assertEqual(resposta.status_code, 200)
        corpo = resposta.json()
        self.assertEqual(corpo["resultados"], [])
        self.assertEqual(corpo["query"], "assunto inexistente")

        # E os dois filtros foram repassados ao buscador, não descartados.
        self.assertEqual(buscador.chamadas[0]["projeto_id"], "SYN-99")
        self.assertEqual(buscador.chamadas[0]["tipo_documento"], "mudancas")

    def test_query_vazia_e_aceita_pelo_schema(self) -> None:
        """CARACTERIZAÇÃO, não garantia.

        `RagSearchRequest.query` é `str` sem `min_length`: string vazia passa e
        a busca é acionada com ela. A ficha de TI-56 já registrava isso como
        caracterização a fazer, e não como comportamento aprovado — vetorizar
        string vazia gasta uma chamada de embedding para nada.
        """
        resposta, buscador = self._buscar({"query": ""})

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(len(buscador.chamadas), 1)
        self.assertEqual(buscador.chamadas[0]["query"], "")

    def test_resultado_atravessa_com_os_metadados_do_indice(self) -> None:
        """Os campos que o RNF12 exige chegam ao cliente, incluindo `chunk_id`."""
        resposta, _ = self._buscar(
            {"query": "avanço", "n_resultados": 1},
            _BuscadorControlado(resultados=[_resultado()]),
        )

        self.assertEqual(resposta.status_code, 200)
        item = resposta.json()["resultados"][0]
        self.assertEqual(
            set(item),
            {"texto", "score", "projeto_id", "tipo_documento", "secao", "arquivo_origem", "chunk_id"},
        )
        self.assertEqual(item["chunk_id"], "chunk-abc", "sem chunk_id o RNF12 fica inverificável")
        self.assertEqual(item["projeto_id"], "SYN-04")


class TestErrosDaBuscaRag(unittest.TestCase):
    """TI-57: cada falha em rodada separada, e a tradução para o contrato."""

    def tearDown(self) -> None:
        limpar_overrides()

    def test_falhas_do_embedding_ou_do_banco_saem_como_500_sem_segredo(self) -> None:
        """Contrato atual: exceção não tratada vira `500 internal_error`.

        Não há fallback nem retry no caminho do RAG, e o caso não inventa
        nenhum dos dois. O que ele garante é que o corpo da resposta não carrega
        a URI do banco, a chave do provedor nem o nome da biblioteca.
        """
        segredo_falso = "postgresql://az1:senha-secreta@host-interno:5432/az1"
        causas = {
            "exceção do provedor de embedding": RuntimeError("falha ao vetorizar (simulada)"),
            "resposta vazia ou malformada": ValueError("resposta sem campo embeddings (simulada)"),
            "vetor de dimensão incompatível": ValueError(
                "expected 1536 dimensions, got 3072 (simulada)"
            ),
            "conexão SQL indisponível": ConnectionError(f"não foi possível conectar a {segredo_falso}"),
        }

        for descricao, erro in causas.items():
            with self.subTest(causa=descricao):
                buscador = _BuscadorControlado(erro=erro)
                http = cliente({get_document_searcher: lambda: buscador})

                resposta = http.post("/api/v1/rag/search", json={"query": "avanço"})

                self.assertEqual(resposta.status_code, 500)
                corpo = resposta.json()
                self.assertEqual(corpo["error"], "internal_error")
                self.assertEqual(corpo["message"], "Erro interno inesperado.")
                for proibido in ("postgresql://", "senha-secreta", "host-interno", "simulada"):
                    self.assertNotIn(proibido, resposta.text)


class TestIngestaoRag(unittest.TestCase):
    """TI-58: extração, fragmentação e o que o pipeline ainda não faz."""

    def setUp(self) -> None:
        self._temporario = tempfile.TemporaryDirectory()
        # O nome da PASTA é o que dá o `projeto_id` (`parsers._projeto_id`),
        # e o prefixo do ARQUIVO é o que dá o `tipo_documento`.
        self.pasta = Path(self._temporario.name) / "SYN-04_Expansao_Linha_6"
        self.pasta.mkdir(parents=True)

    def tearDown(self) -> None:
        self._temporario.cleanup()

    def test_docx_preserva_secao_e_projeto(self) -> None:
        """Título vira `secao`; parágrafos viram trechos do projeto da pasta."""
        caminho = _docx_com_paragrafos_e_tabela(self.pasta)

        extraidos = extrair_docx(caminho)

        self.assertTrue(extraidos)
        for texto in extraidos:
            self.assertEqual(texto.projeto_id, "SYN-04")
            self.assertEqual(texto.tipo_documento, "termo_abertura")
            self.assertEqual(texto.arquivo_origem, "01_TAP.docx")
            self.assertEqual(texto.secao, "Escopo do empreendimento")

        corpo = " ".join(t.texto for t in extraidos)
        self.assertIn("quatro estações", corpo)
        self.assertIn("trinta e seis meses", corpo)

    def test_tabela_de_docx_nao_e_extraida(self) -> None:
        """LACUNA CONHECIDA, registrada em vez de mascarada.

        `parsers.extrair_docx` percorre apenas `doc.paragraphs`; o conteúdo de
        `doc.tables` fica de fora. Num TAP, a tabela costuma ser exatamente onde
        estão marcos e datas — o agente não tem como citá-los, porque eles nunca
        chegaram ao índice.

        A ficha de TI-58 antecipava a lacuna; este caso a torna verificável. Se
        o parser passar a ler tabelas, ele falha e deve ser reescrito como
        verificação positiva.
        """
        caminho = _docx_com_paragrafos_e_tabela(self.pasta)

        extraidos = extrair_docx(caminho)
        corpo = " ".join(t.texto for t in extraidos)

        self.assertNotIn("Conclusão da via permanente", corpo)
        self.assertNotIn("2027-03-31", corpo)

    def test_xlsx_separa_abas_e_ignora_linhas_vazias(self) -> None:
        """Cada aba vira `secao`; linhas sem conteúdo não geram trecho."""
        caminho = _xlsx_com_abas_e_vazios(self.pasta)

        extraidos = extrair_xlsx(caminho)

        secoes = {t.secao for t in extraidos}
        self.assertEqual(secoes, {"Riscos", "Problemas"})
        self.assertTrue(all(t.texto.strip() for t in extraidos))
        self.assertTrue(all(t.tipo_documento == "riscos_problemas" for t in extraidos))

        corpo = " ".join(t.texto for t in extraidos)
        self.assertIn("desapropriação", corpo)
        self.assertIn("perfuração", corpo)

    def test_chunks_nao_misturam_secoes_nem_projetos(self) -> None:
        """A fragmentação respeita a fronteira de arquivo e de seção.

        É o que impede o agente de atribuir a um projeto uma informação que veio
        do documento de outro — o risco que a própria instrução de sistema do
        `gemini_service` passa metade do texto tentando conter.
        """
        docx = _docx_com_paragrafos_e_tabela(self.pasta)
        xlsx = _xlsx_com_abas_e_vazios(self.pasta)

        chunks = chunkar(extrair(docx) + extrair(xlsx))

        self.assertTrue(chunks)
        for chunk in chunks:
            self.assertEqual(chunk.projeto_id, "SYN-04")

        # Nenhum chunk atravessa a fronteira entre dois arquivos...
        por_arquivo = {c.arquivo_origem for c in chunks}
        self.assertEqual(por_arquivo, {"01_TAP.docx", "04_Riscos.xlsx"})

        # ...nem mistura o conteúdo das duas abas da planilha.
        for chunk in chunks:
            if chunk.arquivo_origem == "04_Riscos.xlsx":
                self.assertIn(chunk.secao, {"Riscos", "Problemas"})
                if chunk.secao == "Problemas":
                    self.assertNotIn("desapropriação", chunk.texto)

    def test_reindexar_a_mesma_massa_produz_os_mesmos_identificadores(self) -> None:
        """Repetição idêntica não duplica: o `chunk_id` é função do conteúdo.

        `indexador._chunk_id` é o md5 de `(arquivo, índice, primeiros 60
        caracteres)`, e o `upsert` reaproveita a linha. Este caso verifica a
        metade determinística — os identificadores — sem precisar do banco.
        """
        from rag.indexador import _chunk_id

        caminho = _docx_com_paragrafos_e_tabela(self.pasta)

        primeira = [_chunk_id(c) for c in chunkar(extrair(caminho))]
        segunda = [_chunk_id(c) for c in chunkar(extrair(caminho))]

        self.assertEqual(primeira, segunda)
        self.assertEqual(len(set(primeira)), len(primeira), "identificadores repetidos na mesma massa")

    def test_alterar_o_conteudo_muda_o_identificador_do_chunk(self) -> None:
        """E a outra metade: conteúdo diferente, identificador diferente.

        É o que explica o comentário de `conversa_repository.FonteDaResposta`
        sobre os metadados serem cópia: reindexar um documento editado troca o
        `chunk_id`, e uma trilha que dependesse de junção pelo id perderia a
        ligação com o trecho citado.

        CARACTERIZAÇÃO do outro lado da moeda: nada no pipeline REMOVE o chunk
        antigo. Reindexar uma versão editada deixa os dois no índice, e a busca
        pode devolver o trecho velho. A ficha de TI-58 já antecipava ("revisão
        pode deixar chunks antigos"); confirmado.
        """
        from rag.indexador import _chunk_id

        caminho = _docx_com_paragrafos_e_tabela(self.pasta)
        antes = [_chunk_id(c) for c in chunkar(extrair(caminho))]

        editado = Document()
        editado.add_heading("Escopo do empreendimento", level=1)
        editado.add_paragraph("O projeto contempla a expansão de SEIS estações.")
        editado.save(caminho)

        depois = [_chunk_id(c) for c in chunkar(extrair(caminho))]

        self.assertNotEqual(antes, depois, "editar o conteúdo não mudou o identificador")


@unittest.skipIf(not RAG_DSN, MOTIVO_SEM_BANCO)
class TestBuscaRagPontaAPonta(unittest.TestCase):
    """TI-55: indexação e busca contra `vecs` real.

    Exige `TEST_RAG_DB_URL` e GASTA chamadas reais de embedding: os vetores de
    DOCUMENTO não estão em fita, e a fita de embedding que existe é de consulta
    — tarefa diferente, vetor diferente. Por isso o opt-in é explícito, em vez
    de a suíte tentar e falhar.
    """

    @classmethod
    def setUpClass(cls) -> None:
        # `indexador._db_url` lê SUPABASE_DB_URL; apontá-la para a base de teste
        # é o que impede a indexação sintética de cair no índice de produção.
        cls._anterior = os.environ.get("SUPABASE_DB_URL")
        os.environ["SUPABASE_DB_URL"] = RAG_DSN

        from rag import indexador

        indexador._cliente.cache_clear()  # noqa: SLF001 - o cache guarda a URL antiga

    @classmethod
    def tearDownClass(cls) -> None:
        from rag import indexador

        indexador._cliente.cache_clear()  # noqa: SLF001
        if cls._anterior is None:
            os.environ.pop("SUPABASE_DB_URL", None)
        else:
            os.environ["SUPABASE_DB_URL"] = cls._anterior

    def setUp(self) -> None:
        self._temporario = tempfile.TemporaryDirectory()
        self.raiz = Path(self._temporario.name)

    def tearDown(self) -> None:
        limpar_overrides()
        self._temporario.cleanup()

    def test_busca_com_filtro_combinado_respeita_projeto_e_tipo(self) -> None:
        """Dois documentos sintéticos de projetos distintos, e o filtro separa."""
        from rag.embedder import vetorizar_documentos
        from rag.indexador import indexar
        from rag.retriever import buscar

        chunks: list = []
        for projeto in ("SYN-04_Primeiro", "SYN-07_Segundo"):
            pasta = self.raiz / projeto
            pasta.mkdir()
            chunks.extend(chunkar(extrair(_docx_com_paragrafos_e_tabela(pasta))))

        self.assertTrue(chunks, "a massa sintética não produziu chunk algum")
        indexar(chunks, vetorizar_documentos([c.texto for c in chunks]))

        resultados = buscar(
            "expansão de estações",
            n_resultados=1,
            projeto_id="SYN-04",
            tipo_documento="termo_abertura",
        )

        self.assertLessEqual(len(resultados), 1)
        for item in resultados:
            self.assertEqual(item.projeto_id, "SYN-04")
            self.assertEqual(item.tipo_documento, "termo_abertura")
            self.assertEqual(item.arquivo_origem, "01_TAP.docx")
            self.assertTrue(item.chunk_id)


if __name__ == "__main__":
    unittest.main()
